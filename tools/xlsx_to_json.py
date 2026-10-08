#!/usr/bin/env python3
"""Dumps an .xlsx as {sheetName: rows[][]} JSON (used by tests to feed the parser)."""
import json, sys
from datetime import date, datetime
from openpyxl import load_workbook

wb = load_workbook(sys.argv[1], data_only=True)
def conv(v):
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    return v
print(json.dumps({ws.title: [[conv(c) for c in row] for row in ws.iter_rows(values_only=True)] for ws in wb.worksheets}))
